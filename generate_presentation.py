"""
Générateur de présentation PowerPoint - HARDTEC AIOps
20 slides pour soutenance PFA - ENSA Berrechid
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ============================================================
# PALETTE DE COULEURS
# ============================================================
BLEU_NUIT = RGBColor(0x0A, 0x16, 0x28)
BLEU = RGBColor(0x25, 0x63, 0xEB)
TURQUOISE = RGBColor(0x00, 0xD4, 0xAA)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
GRIS_CLAIR = RGBColor(0xF5, 0xF7, 0xFA)
GRIS_MOYEN = RGBColor(0x6B, 0x72, 0x80)
GRIS_FONCE = RGBColor(0x1F, 0x29, 0x37)
ORANGE = RGBColor(0xF5, 0x9E, 0x0B)
VERT = RGBColor(0x10, 0xB9, 0x81)
ROUGE = RGBColor(0xEF, 0x44, 0x44)

# ============================================================
# INITIALISATION
# ============================================================
prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)
BLANK = prs.slide_layouts[6]


# ============================================================
# FONCTIONS UTILITAIRES
# ============================================================
def add_slide():
    return prs.slides.add_slide(BLANK)


def add_rect(slide, x, y, w, h, fill=BLANC, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, size=18, bold=False,
             color=GRIS_FONCE, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return tb


def add_title(slide, text):
    """Ajoute un titre standard en haut de slide."""
    add_rect(slide, 0, 0, prs.slide_width, Inches(1.0), fill=BLEU_NUIT)
    add_text(slide, Inches(0.5), Inches(0.2), Inches(15), Inches(0.7),
             text, size=28, bold=True, color=BLANC)
    # Ligne turquoise de séparation
    add_rect(slide, 0, Inches(1.0), prs.slide_width, Inches(0.05),
             fill=TURQUOISE)


def add_bullet_box(slide, x, y, w, h, title, bullets,
                   title_color=BLEU_NUIT, box_fill=GRIS_CLAIR,
                   border_color=TURQUOISE):
    add_rect(slide, x, y, w, h, fill=box_fill, line=border_color, radius=True)
    add_text(slide, x + Inches(0.2), y + Inches(0.15),
             w - Inches(0.4), Inches(0.5),
             title, size=16, bold=True, color=title_color)
    body = "\n".join([f"• {b}" for b in bullets])
    add_text(slide, x + Inches(0.2), y + Inches(0.7),
             w - Inches(0.4), h - Inches(0.8),
             body, size=12, color=GRIS_FONCE)


# ============================================================
# SLIDE 1 — PAGE DE GARDE
# ============================================================
s = add_slide()
add_rect(s, 0, 0, prs.slide_width, prs.slide_height, fill=BLEU_NUIT)
add_text(s, Inches(0.5), Inches(0.4), Inches(15), Inches(0.5),
         "ENSA BERRECHID  |  Big Data & Systèmes d'Information",
         size=14, color=TURQUOISE, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.5), Inches(2.2), Inches(15), Inches(1.2),
         "HARDTEC AIOps", size=60, bold=True, color=BLANC,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(0.5), Inches(3.4), Inches(15), Inches(0.7),
         "Intelligent Support Platform", size=32, color=TURQUOISE,
         align=PP_ALIGN.CENTER)

add_rect(s, Inches(5), Inches(4.3), Inches(6), Inches(0.03), fill=TURQUOISE)

add_text(s, Inches(1), Inches(4.6), Inches(14), Inches(1.0),
         "Conception et développement d'une plateforme AIOps intelligente\n"
         "pour la détection, la prédiction et l'assistance aux incidents informatiques",
         size=16, color=BLANC, align=PP_ALIGN.CENTER)

add_text(s, Inches(1), Inches(6.3), Inches(14), Inches(0.4),
         "Projet de Fin d'Année 2025–2026", size=14, color=GRIS_MOYEN,
         align=PP_ALIGN.CENTER)

add_text(s, Inches(1), Inches(7.0), Inches(14), Inches(1.5),
         "Entreprise : HARDTEC MAROC\n"
         "Encadrant professionnel : M. FADIL BADR\n"
         "Encadrant académique : Pr. Ahmed Nafidi\n"
         "Présenté par : [Nom et Prénom]",
         size=13, color=BLANC, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 2 — CONTEXTE
# ============================================================
s = add_slide()
add_title(s, "Contexte et problématique")

# 3 blocs pipeline
boxes = [
    ("TICKETS", ["Volume important", "Hétérogènes", "Multilingues", "Non structurés"]),
    ("ANALYSE", ["Manuelle", "Lente", "Coûteuse", "Répétitive"]),
    ("DÉCISION", ["Routage", "Erreurs possibles", "Mauvaise affectation"]),
]
x0 = Inches(0.8)
for i, (title, items) in enumerate(boxes):
    x = x0 + i * Inches(5)
    add_bullet_box(s, x, Inches(1.5), Inches(4.5), Inches(3.2), title, items)
    if i < 2:
        add_text(s, x + Inches(4.5), Inches(2.8), Inches(0.5), Inches(0.5),
                 "→", size=32, bold=True, color=TURQUOISE, align=PP_ALIGN.CENTER)

# Problématique
add_rect(s, Inches(0.8), Inches(5.2), Inches(14.4), Inches(1.8),
         fill=GRIS_CLAIR, line=TURQUOISE, radius=True)
add_text(s, Inches(1.2), Inches(5.4), Inches(13.6), Inches(0.4),
         "PROBLÉMATIQUE CENTRALE", size=14, bold=True, color=BLEU_NUIT)
add_text(s, Inches(1.2), Inches(5.9), Inches(13.6), Inches(1.0),
         "Comment exploiter les données IT, le Machine Learning et l'IA\n"
         "pour mieux classifier les tickets, anticiper les risques et assister le support ?",
         size=16, color=GRIS_FONCE)

add_text(s, Inches(0.8), Inches(7.4), Inches(14.4), Inches(0.5),
         "💡  L'objectif : ASSISTER l'opérateur, pas le REMPLACER",
         size=16, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 3 — OBJECTIFS
# ============================================================
s = add_slide()
add_title(s, "Objectifs du projet")

cards = [
    ("1. DATA ENGINEERING", "Structurer et préparer les données", BLEU),
    ("2. TICKET INTELLIGENCE", "Classifier et prédire l'affectation", TURQUOISE),
    ("3. AIOps", "Expérimenter la prédiction du risque", BLEU),
    ("4. INTELLIGENT SUPPORT", "RAG + Agent IA + Recherche sémantique", TURQUOISE),
]
positions = [
    (Inches(0.8), Inches(1.5)),
    (Inches(8.2), Inches(1.5)),
    (Inches(0.8), Inches(4.3)),
    (Inches(8.2), Inches(4.3)),
]
for (title, desc, color), (x, y) in zip(cards, positions):
    add_rect(s, x, y, Inches(7.0), Inches(2.4), fill=GRIS_CLAIR,
             line=color, radius=True)
    add_rect(s, x, y, Inches(7.0), Inches(0.6), fill=color, radius=True)
    add_text(s, x + Inches(0.2), y + Inches(0.1), Inches(6.6), Inches(0.5),
             title, size=18, bold=True, color=BLANC)
    add_text(s, x + Inches(0.3), y + Inches(0.9), Inches(6.4), Inches(1.3),
             desc, size=16, color=GRIS_FONCE)

add_rect(s, Inches(0.8), Inches(7.0), Inches(14.4), Inches(0.9),
         fill=ORANGE, radius=True)
add_text(s, Inches(1.0), Inches(7.15), Inches(14.0), Inches(0.6),
         "🎯  Principe directeur : HUMAN-IN-THE-LOOP — L'IA propose, l'opérateur valide",
         size=15, bold=True, color=BLANC, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 4 — PÉRIMÈTRE
# ============================================================
s = add_slide()
add_title(s, "Périmètre réalisé")

# Implémenté
add_rect(s, Inches(0.8), Inches(1.5), Inches(4.6), Inches(5.5),
         fill=RGBColor(0xEC, 0xFD, 0xF5), line=VERT, radius=True)
add_text(s, Inches(1.0), Inches(1.7), Inches(4.2), Inches(0.5),
         "✅  IMPLÉMENTÉ", size=18, bold=True, color=VERT)
impl = ["Pipeline classification", "API FastAPI (9 endpoints)",
        "Interface Streamlit (5 pages)", "Historique SQLite",
        "Modèles Type/Queue/Priority", "4 DAGs Airflow",
        "Docker Compose", "MLflow tracking"]
add_text(s, Inches(1.0), Inches(2.4), Inches(4.2), Inches(4.5),
         "\n".join([f"• {i}" for i in impl]), size=13, color=GRIS_FONCE)

# Expérimental
add_rect(s, Inches(5.7), Inches(1.5), Inches(4.6), Inches(5.5),
         fill=RGBColor(0xFF, 0xFB, 0xEB), line=ORANGE, radius=True)
add_text(s, Inches(5.9), Inches(1.7), Inches(4.2), Inches(0.5),
         "⚠️  EXPÉRIMENTAL", size=18, bold=True, color=ORANGE)
exp = ["Prédiction AIOps", "(performances faibles)",
       "Artefacts RAG absents", "(tickets.index, documents.pkl)",
       "Validation end-to-end", "non démontrée",
       "Métriques Queue/Priority", "à compléter"]
add_text(s, Inches(5.9), Inches(2.4), Inches(4.2), Inches(4.5),
         "\n".join([f"• {i}" for i in exp]), size=13, color=GRIS_FONCE)

# Perspectives
add_rect(s, Inches(10.6), Inches(1.5), Inches(4.6), Inches(5.5),
         fill=RGBColor(0xEF, 0xF6, 0xFF), line=BLEU, radius=True)
add_text(s, Inches(10.8), Inches(1.7), Inches(4.2), Inches(0.5),
         "🔮  PERSPECTIVES", size=18, bold=True, color=BLEU)
persp = ["Monitoring temps réel", "Amélioration des modèles",
         "Kafka (streaming)", "Industrialisation complète",
         "Observabilité (Grafana)", "Kubernetes",
         "Calibration des probabilités", "Authentification API"]
add_text(s, Inches(10.8), Inches(2.4), Inches(4.2), Inches(4.5),
         "\n".join([f"• {i}" for i in persp]), size=13, color=GRIS_FONCE)


# ============================================================
# SLIDE 5 — ARCHITECTURE
# ============================================================
s = add_slide()
add_title(s, "Architecture globale de la plateforme")

def arch_box(x, y, w, h, text, fill, text_color=BLANC, size=12):
    add_rect(s, x, y, w, h, fill=fill, radius=True)
    add_text(s, x + Inches(0.1), y + Inches(0.15), w - Inches(0.2), h - Inches(0.3),
             text, size=size, bold=True, color=text_color, align=PP_ALIGN.CENTER)

def arrow(x, y, w=0.4, h=0.3):
    add_text(s, x, y, Inches(w), Inches(h), "▼",
             size=18, bold=True, color=TURQUOISE, align=PP_ALIGN.CENTER)

# Sources
arch_box(Inches(5.5), Inches(1.2), Inches(5), Inches(0.6),
         "SOURCES : Tickets IT  |  Télémétrie MicroSS", BLEU)
arrow(Inches(7.9), Inches(1.85))

# Data eng
arch_box(Inches(5.5), Inches(2.2), Inches(5), Inches(0.6),
         "DATA ENGINEERING : Snowflake | dbt | DuckDB | SQLite", BLEU)
arrow(Inches(7.9), Inches(2.85))

# Ticket & AIOps
arch_box(Inches(2.5), Inches(3.3), Inches(4.5), Inches(1.1),
         "TICKET INTELLIGENCE\nType → Queue → Priority", TURQUOISE)
arch_box(Inches(9.0), Inches(3.3), Inches(4.5), Inches(1.1),
         "AIOps\nMicroSS / XGBoost", ORANGE)
arrow(Inches(7.9), Inches(4.45))

# Intelligent Support
arch_box(Inches(5.5), Inches(4.9), Inches(5), Inches(0.7),
         "INTELLIGENT SUPPORT : RAG (FAISS) + Agent IA", TURQUOISE)
arrow(Inches(7.9), Inches(5.65))

# API & Interface
arch_box(Inches(5.5), Inches(6.0), Inches(5), Inches(0.6),
         "FastAPI  +  Streamlit", BLEU)
arrow(Inches(7.9), Inches(6.65))

# Historique
arch_box(Inches(5.5), Inches(7.0), Inches(5), Inches(0.6),
         "Historique SQLite", GRIS_FONCE)

# Colonne latérale outils
add_text(s, Inches(0.3), Inches(5.5), Inches(2), Inches(2),
         "Airflow\nOrchestration\n\nDocker\nConteneurisation\n\nMLflow\nSuivi exp.",
         size=11, color=GRIS_MOYEN)


# ============================================================
# SLIDE 6 — DATA ENGINEERING
# ============================================================
s = add_slide()
add_title(s, "Data Engineering et préparation des données")

steps = ["DONNÉES\nBRUTES", "INGESTION", "TRANSFORMATION", "VALIDATION",
         "DONNÉES\nEXPLOITABLES"]
x0 = Inches(0.8)
for i, step in enumerate(steps):
    x = x0 + i * Inches(3.0)
    add_rect(s, x, Inches(1.6), Inches(2.5), Inches(1.2),
             fill=BLEU if i % 2 == 0 else TURQUOISE, radius=True)
    add_text(s, x + Inches(0.1), Inches(1.85), Inches(2.3), Inches(0.9),
             step, size=12, bold=True, color=BLANC, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        add_text(s, x + Inches(2.5), Inches(1.95), Inches(0.5), Inches(0.5),
                 "▶", size=20, bold=True, color=TURQUOISE, align=PP_ALIGN.CENTER)

add_rect(s, Inches(0.8), Inches(3.3), Inches(14.4), Inches(0.7),
         fill=GRIS_CLAIR, line=TURQUOISE, radius=True)
add_text(s, Inches(1.0), Inches(3.45), Inches(14.0), Inches(0.5),
         "📊  20 000 tickets prétraités  —  tickets_clean.csv",
         size=16, bold=True, color=BLEU_NUIT, align=PP_ALIGN.CENTER)

# Technologies
add_rect(s, Inches(0.8), Inches(4.3), Inches(14.4), Inches(3.5),
         fill=GRIS_CLAIR, line=TURQUOISE, radius=True)
add_text(s, Inches(1.1), Inches(4.5), Inches(13.8), Inches(0.5),
         "TECHNOLOGIES ET RÔLES", size=18, bold=True, color=BLEU_NUIT)

techs = [
    ("❄️  Snowflake", "Data Warehouse — stockage centralisé"),
    ("🔧  dbt", "Transformation SQL — Bronze / Silver / Gold"),
    ("🦆  DuckDB", "Analyse locale OLAP"),
    ("💾  SQLite", "Persistance transactionnelle"),
]
for i, (name, role) in enumerate(techs):
    y = Inches(5.1) + i * Inches(0.65)
    add_text(s, Inches(1.3), y, Inches(4), Inches(0.5),
             name, size=15, bold=True, color=BLEU)
    add_text(s, Inches(5.5), y, Inches(9.5), Inches(0.5),
             role, size=15, color=GRIS_FONCE)


# ============================================================
# SLIDE 7 — MACHINE LEARNING
# ============================================================
s = add_slide()
add_title(s, "Classification intelligente des tickets")

# Ticket entrant
add_rect(s, Inches(3), Inches(1.3), Inches(10), Inches(0.9),
         fill=BLEU_NUIT, radius=True)
add_text(s, Inches(3.2), Inches(1.45), Inches(9.6), Inches(0.6),
         "📩  TICKET UTILISATEUR : « Le serveur est complètement down... »",
         size=15, bold=True, color=BLANC, align=PP_ALIGN.CENTER)

steps_ml = [
    ("ÉTAPE 1 : TYPE", "ticket_type_model.pkl", "TF-IDF + LinearSVC", "Incident | Request | Problem | Change"),
    ("ÉTAPE 2 : QUEUE", "ticket_queue_model_v4.pkl", "TF-IDF + OneHot(Type) + LinearSVC", "Technical Support | IT Support | Billing | ..."),
    ("ÉTAPE 3 : PRIORITY", "ticket_priority_model_v3.pkl", "TF-IDF + OneHot(Type, Queue) + LinearSVC", "High | Medium | Low"),
]
for i, (title, model, method, output) in enumerate(steps_ml):
    y = Inches(2.5) + i * Inches(1.55)
    add_rect(s, Inches(2), y, Inches(12), Inches(1.35),
             fill=GRIS_CLAIR, line=TURQUOISE, radius=True)
    add_text(s, Inches(2.3), y + Inches(0.1), Inches(11.4), Inches(0.4),
             title, size=15, bold=True, color=BLEU_NUIT)
    add_text(s, Inches(2.3), y + Inches(0.5), Inches(6), Inches(0.4),
             f"Modèle : {model}", size=12, color=BLEU)
    add_text(s, Inches(2.3), y + Inches(0.85), Inches(11.4), Inches(0.4),
             f"{method}   →   {output}", size=11, color=GRIS_FONCE)
    if i < len(steps_ml) - 1:
        add_text(s, Inches(7.8), y + Inches(1.3), Inches(0.5), Inches(0.4),
                 "▼", size=20, bold=True, color=TURQUOISE, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.8), Inches(7.5), Inches(14.4), Inches(0.5),
         "⚠️  Pipeline en cascade : une erreur en amont peut se propager vers les étapes suivantes",
         size=14, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 8 — RÉSULTATS ML
# ============================================================
s = add_slide()
add_title(s, "Résultats — Classification du type")

# Deux grosses métriques
add_rect(s, Inches(2), Inches(1.5), Inches(5.5), Inches(2.2),
         fill=RGBColor(0xEC, 0xFD, 0xF5), line=VERT, radius=True)
add_text(s, Inches(2.2), Inches(1.7), Inches(5.1), Inches(0.5),
         "ACCURACY", size=18, bold=True, color=VERT, align=PP_ALIGN.CENTER)
add_text(s, Inches(2.2), Inches(2.3), Inches(5.1), Inches(1.2),
         "81,98 %", size=54, bold=True, color=GRIS_FONCE, align=PP_ALIGN.CENTER)

add_rect(s, Inches(8.5), Inches(1.5), Inches(5.5), Inches(2.2),
         fill=RGBColor(0xEC, 0xFD, 0xF5), line=VERT, radius=True)
add_text(s, Inches(8.7), Inches(1.7), Inches(5.1), Inches(0.5),
         "MACRO-F1", size=18, bold=True, color=VERT, align=PP_ALIGN.CENTER)
add_text(s, Inches(8.7), Inches(2.3), Inches(5.1), Inches(1.2),
         "81,80 %", size=54, bold=True, color=GRIS_FONCE, align=PP_ALIGN.CENTER)

add_text(s, Inches(2), Inches(3.9), Inches(12), Inches(0.5),
         "📊  Évaluation sur 4 000 exemples de test",
         size=15, bold=True, color=BLEU_NUIT, align=PP_ALIGN.CENTER)

# F1 par classe
add_text(s, Inches(1), Inches(4.5), Inches(6), Inches(0.4),
         "F1-SCORE PAR CLASSE", size=14, bold=True, color=BLEU_NUIT)
classes = [("Request", 97.98, VERT), ("Change", 94.42, VERT),
           ("Incident", 79.93, ORANGE), ("Problem", 54.86, ROUGE)]
for i, (name, score, color) in enumerate(classes):
    y = Inches(5.0) + i * Inches(0.55)
    add_text(s, Inches(1), y, Inches(1.5), Inches(0.4),
             name, size=13, bold=True, color=GRIS_FONCE)
    bar_width = Inches(score / 100 * 9)
    add_rect(s, Inches(2.5), y + Inches(0.05), bar_width, Inches(0.35),
             fill=color)
    add_text(s, Inches(2.5) + bar_width + Inches(0.1), y, Inches(1.5), Inches(0.4),
             f"{score:.2f} %", size=12, bold=True, color=GRIS_FONCE)

add_text(s, Inches(1), Inches(7.6), Inches(14), Inches(0.6),
         "✅  Base solide pour automatiser l'orientation    |    "
         "⚠️  Classe Problem plus difficile à distinguer",
         size=13, bold=True, color=BLEU_NUIT, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 9 — AIOPS
# ============================================================
s = add_slide()
add_title(s, "AIOps — Prédiction du risque d'incident")

# Badge expérimental
add_rect(s, Inches(13.5), Inches(0.15), Inches(2.2), Inches(0.6),
         fill=ORANGE, radius=True)
add_text(s, Inches(13.6), Inches(0.25), Inches(2.0), Inches(0.4),
         "EXPÉRIMENTAL", size=12, bold=True, color=BLANC, align=PP_ALIGN.CENTER)

steps_aiops = [
    ("115 métriques de base", "CPU, mémoire, I/O réseau, taux d'erreurs", BLEU),
    ("Feature Engineering", "→ 1 380 features temporelles (lags, deltas, % var., rolling stats)", TURQUOISE),
    ("MicroSS / XGBoost V3.1", "→ Probabilité de risque d'incident", BLEU),
    ("Seuil de décision : 0,35", "→ Alerte risque  |  État normal", ORANGE),
]
for i, (title, sub, color) in enumerate(steps_aiops):
    y = Inches(1.4) + i * Inches(1.3)
    add_rect(s, Inches(3), y, Inches(10), Inches(1.05),
             fill=GRIS_CLAIR, line=color, radius=True)
    add_text(s, Inches(3.3), y + Inches(0.1), Inches(9.4), Inches(0.5),
             title, size=17, bold=True, color=color)
    add_text(s, Inches(3.3), y + Inches(0.55), Inches(9.4), Inches(0.4),
             sub, size=13, color=GRIS_FONCE)
    if i < len(steps_aiops) - 1:
        add_text(s, Inches(7.8), y + Inches(1.0), Inches(0.5), Inches(0.3),
                 "▼", size=18, bold=True, color=TURQUOISE, align=PP_ALIGN.CENTER)

# Infos clés
add_rect(s, Inches(1), Inches(7.0), Inches(6.5), Inches(0.7),
         fill=BLEU, radius=True)
add_text(s, Inches(1.2), Inches(7.15), Inches(6.1), Inches(0.5),
         "Granularité : 5 minutes", size=15, bold=True, color=BLANC,
         align=PP_ALIGN.CENTER)

add_rect(s, Inches(8.5), Inches(7.0), Inches(6.5), Inches(0.7),
         fill=BLEU, radius=True)
add_text(s, Inches(8.7), Inches(7.15), Inches(6.1), Inches(0.5),
         "Horizon : 15 minutes", size=15, bold=True, color=BLANC,
         align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 10 — RÉSULTATS AIOPS
# ============================================================
s = add_slide()
add_title(s, "Résultats — Modèle AIOps")

metrics = [
    ("PRECISION", "3,20 %", ROUGE),
    ("RECALL", "42,00 %", ORANGE),
    ("F1-SCORE", "5,95 %", ROUGE),
]
for i, (name, value, color) in enumerate(metrics):
    x = Inches(1) + i * Inches(4.8)
    add_rect(s, x, Inches(1.5), Inches(4.2), Inches(2.0),
             fill=GRIS_CLAIR, line=color, radius=True)
    add_text(s, x + Inches(0.2), Inches(1.7), Inches(3.8), Inches(0.4),
             name, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.3), Inches(3.8), Inches(1.0),
             value, size=42, bold=True, color=GRIS_FONCE, align=PP_ALIGN.CENTER)

metrics2 = [("PR-AUC", "3,27 %", ROUGE), ("ROC-AUC", "55,40 %", ORANGE)]
for i, (name, value, color) in enumerate(metrics2):
    x = Inches(2.5) + i * Inches(6.5)
    add_rect(s, x, Inches(3.8), Inches(5.5), Inches(1.6),
             fill=GRIS_CLAIR, line=color, radius=True)
    add_text(s, x + Inches(0.2), Inches(3.95), Inches(5.1), Inches(0.4),
             name, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(4.4), Inches(5.1), Inches(0.9),
             value, size=38, bold=True, color=GRIS_FONCE, align=PP_ALIGN.CENTER)

# Interprétation
add_rect(s, Inches(0.8), Inches(5.8), Inches(14.4), Inches(2.2),
         fill=RGBColor(0xFF, 0xFB, 0xEB), line=ORANGE, radius=True)
add_text(s, Inches(1.1), Inches(5.95), Inches(13.8), Inches(0.5),
         "INTERPRÉTATION HONNÊTE", size=16, bold=True, color=ORANGE)
add_text(s, Inches(1.1), Inches(6.45), Inches(13.8), Inches(1.5),
         "•  Recall 42 % → détecte une partie des incidents\n"
         "•  Precision 3,2 % → génère énormément de faux positifs\n"
         "•  ROC-AUC 0,554 → proche du hasard\n"
         "⚠️  Modèle expérimental nécessitant des améliorations avant tout usage opérationnel",
         size=13, color=GRIS_FONCE)


# ============================================================
# SLIDE 11 — RAG
# ============================================================
s = add_slide()
add_title(s, "Recherche sémantique et RAG")

rag_steps = [
    ("REQUÊTE UTILISATEUR", "", BLEU),
    ("SENTENCE TRANSFORMER", "paraphrase-multilingual-MiniLM-L12-v2", TURQUOISE),
    ("EMBEDDING", "Vecteur dense float32 + Normalisation L2", BLEU),
    ("FAISS", "Recherche Top-K (k = 5)", TURQUOISE),
    ("TICKETS SIMILAIRES", "Métadonnées + solutions", BLEU),
    ("ASSISTANCE", "Contexte pour l'opérateur", TURQUOISE),
]
for i, (title, sub, color) in enumerate(rag_steps):
    y = Inches(1.3) + i * Inches(0.95)
    add_rect(s, Inches(3.5), y, Inches(9), Inches(0.75),
             fill=color, radius=True)
    text = title if not sub else f"{title}  —  {sub}"
    add_text(s, Inches(3.7), y + Inches(0.18), Inches(8.6), Inches(0.5),
             text, size=13, bold=True, color=BLANC, align=PP_ALIGN.CENTER)
    if i < len(rag_steps) - 1:
        add_text(s, Inches(7.8), y + Inches(0.75), Inches(0.4), Inches(0.2),
                 "▼", size=12, bold=True, color=GRIS_MOYEN, align=PP_ALIGN.CENTER)

# Avertissement
add_rect(s, Inches(1), Inches(7.2), Inches(14), Inches(1.0),
         fill=RGBColor(0xFF, 0xFB, 0xEB), line=ORANGE, radius=True)
add_text(s, Inches(1.2), Inches(7.3), Inches(13.6), Inches(0.4),
         "⚠️  À VALIDER — Artefacts absents de l'archive finale :",
         size=13, bold=True, color=ORANGE)
add_text(s, Inches(1.2), Inches(7.65), Inches(13.6), Inches(0.5),
         "models/rag/tickets.index   •   models/rag/documents.pkl",
         size=13, color=GRIS_FONCE)


# ============================================================
# SLIDE 12 — AGENT IA
# ============================================================
s = add_slide()
add_title(s, "Agent IA — Assistance à l'analyse des incidents")

agent_steps = [
    ("📩  TICKET / INCIDENT", "", BLEU_NUIT),
    ("ANALYSE", "Prédiction ML + Recherche RAG + Score AIOps", TURQUOISE),
    ("CONTEXTE DISPONIBLE", "Type, Queue, Priority, tickets similaires, risque", BLEU),
    ("RECOMMANDATION / ASSISTANCE", "Plan d'action proposé à l'opérateur", TURQUOISE),
    ("✋  VALIDATION HUMAINE", "HUMAN-IN-THE-LOOP — L'agent propose, l'opérateur valide", VERT),
]
for i, (title, sub, color) in enumerate(agent_steps):
    y = Inches(1.3) + i * Inches(1.15)
    add_rect(s, Inches(2.5), y, Inches(11), Inches(0.95),
             fill=color, radius=True)
    add_text(s, Inches(2.7), y + Inches(0.1), Inches(10.6), Inches(0.5),
             title, size=15, bold=True, color=BLANC, align=PP_ALIGN.CENTER)
    if sub:
        add_text(s, Inches(2.7), y + Inches(0.5), Inches(10.6), Inches(0.4),
                 sub, size=11, color=BLANC, align=PP_ALIGN.CENTER)
    if i < len(agent_steps) - 1:
        add_text(s, Inches(7.8), y + Inches(0.95), Inches(0.5), Inches(0.2),
                 "▼", size=14, bold=True, color=GRIS_MOYEN, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.8), Inches(7.6), Inches(14.4), Inches(0.5),
         "📡  Endpoint : POST /agent/analyze",
         size=15, bold=True, color=BLEU, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 13 — API & STREAMLIT
# ============================================================
s = add_slide()
add_title(s, "API et interface utilisateur")

# FastAPI
add_rect(s, Inches(0.8), Inches(1.4), Inches(7), Inches(6.8),
         fill=GRIS_CLAIR, line=BLEU, radius=True)
add_rect(s, Inches(0.8), Inches(1.4), Inches(7), Inches(0.6),
         fill=BLEU, radius=True)
add_text(s, Inches(1.0), Inches(1.5), Inches(6.6), Inches(0.5),
         "FASTAPI — Backend REST", size=16, bold=True, color=BLANC)

endpoints = [
    ("GET", "/health", "Vérifier l'état des services"),
    ("POST", "/predict", "Classifier un ticket"),
    ("POST", "/aiops/predict-risk", "Prédire le risque AIOps"),
    ("POST", "/rag/search", "Rechercher tickets similaires"),
    ("POST", "/support/analyze", "Analyse ML + RAG + AIOps"),
    ("GET", "/history", "Consulter l'historique"),
    ("GET", "/history/{id}", "Récupérer une prédiction"),
    ("GET", "/stats", "Statistiques agrégées"),
]
for i, (method, path, desc) in enumerate(endpoints):
    y = Inches(2.2) + i * Inches(0.6)
    color = VERT if method == "GET" else ORANGE
    add_text(s, Inches(1.0), y, Inches(1.0), Inches(0.4),
             method, size=11, bold=True, color=color)
    add_text(s, Inches(2.0), y, Inches(3.5), Inches(0.4),
             path, size=12, bold=True, color=BLEU_NUIT)
    add_text(s, Inches(5.5), y, Inches(2.2), Inches(0.4),
             desc, size=10, color=GRIS_FONCE)

add_text(s, Inches(1.0), Inches(7.2), Inches(6.6), Inches(0.5),
         "✅  Pydantic  |  ✅  Erreurs 400/404/500/503  |  Port 8000",
         size=11, color=GRIS_FONCE)

# Streamlit
add_rect(s, Inches(8.2), Inches(1.4), Inches(7), Inches(6.8),
         fill=GRIS_CLAIR, line=TURQUOISE, radius=True)
add_rect(s, Inches(8.2), Inches(1.4), Inches(7), Inches(0.6),
         fill=TURQUOISE, radius=True)
add_text(s, Inches(8.4), Inches(1.5), Inches(6.6), Inches(0.5),
         "STREAMLIT — Interface utilisateur", size=16, bold=True, color=BLANC)

# Placeholder capture
add_rect(s, Inches(8.5), Inches(2.2), Inches(6.4), Inches(4.0),
         fill=RGBColor(0xE5, 0xE7, 0xEB), line=GRIS_MOYEN, radius=True)
add_text(s, Inches(8.5), Inches(3.9), Inches(6.4), Inches(0.5),
         "📸  [CAPTURE À INSÉRER]", size=16, bold=True,
         color=GRIS_MOYEN, align=PP_ALIGN.CENTER)

pages = ["Dashboard", "AI Ticket Analyzer", "Incident Prediction",
         "Intelligent Support", "Analytics"]
add_text(s, Inches(8.5), Inches(6.4), Inches(6.4), Inches(0.4),
         "5 pages :", size=12, bold=True, color=BLEU_NUIT)
add_text(s, Inches(8.5), Inches(6.75), Inches(6.4), Inches(1.4),
         "  •  ".join(pages), size=11, color=GRIS_FONCE)


# ============================================================
# SLIDE 14 — AIRFLOW
# ============================================================
s = add_slide()
add_title(s, "Orchestration des pipelines avec Airflow")

# Airflow central
add_rect(s, Inches(5.5), Inches(1.3), Inches(5), Inches(0.9),
         fill=BLEU_NUIT, radius=True)
add_text(s, Inches(5.7), Inches(1.45), Inches(4.6), Inches(0.6),
         "AIRFLOW — Orchestrateur", size=18, bold=True, color=BLANC,
         align=PP_ALIGN.CENTER)

add_text(s, Inches(7.8), Inches(2.2), Inches(0.5), Inches(0.4),
         "▼", size=22, bold=True, color=TURQUOISE, align=PP_ALIGN.CENTER)

# 3 pipelines
dags = [
    ("hardtec_ticket_pipeline", BLEU),
    ("hardtec_aiops_pipeline", ORANGE),
    ("hardtec_rag_pipeline", TURQUOISE),
]
for i, (name, color) in enumerate(dags):
    x = Inches(0.8) + i * Inches(5.0)
    add_rect(s, x, Inches(2.9), Inches(4.5), Inches(1.2),
             fill=color, radius=True)
    add_text(s, x + Inches(0.2), Inches(3.05), Inches(4.1), Inches(0.4),
             f"DAG {i+1}", size=11, bold=True, color=BLANC,
             align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(3.4), Inches(4.1), Inches(0.6),
             name, size=13, bold=True, color=BLANC, align=PP_ALIGN.CENTER)

# Convergence
for i in range(3):
    x = Inches(3.0) + i * Inches(5.0)
    add_text(s, x, Inches(4.15), Inches(0.5), Inches(0.3),
             "▼", size=14, bold=True, color=GRIS_MOYEN, align=PP_ALIGN.CENTER)

# Complete pipeline
add_rect(s, Inches(2.5), Inches(4.6), Inches(11), Inches(1.6),
         fill=GRIS_CLAIR, line=BLEU_NUIT, radius=True)
add_text(s, Inches(2.7), Inches(4.75), Inches(10.6), Inches(0.5),
         "hardtec_complete_pipeline", size=18, bold=True,
         color=BLEU_NUIT, align=PP_ALIGN.CENTER)
add_text(s, Inches(2.7), Inches(5.35), Inches(10.6), Inches(0.7),
         "Validation → Classification → RAG → AIOps → Synthèse",
         size=14, color=GRIS_FONCE, align=PP_ALIGN.CENTER)

# Infos
add_rect(s, Inches(0.8), Inches(6.6), Inches(7), Inches(0.8),
         fill=BLEU, radius=True)
add_text(s, Inches(1.0), Inches(6.75), Inches(6.6), Inches(0.5),
         "📅  Fréquence : @daily (quotidien)", size=14, bold=True,
         color=BLANC, align=PP_ALIGN.CENTER)

add_rect(s, Inches(8.2), Inches(6.6), Inches(7), Inches(0.8),
         fill=BLEU, radius=True)
add_text(s, Inches(8.4), Inches(6.75), Inches(6.6), Inches(0.5),
         "🎯  4 DAGs configurés", size=14, bold=True,
         color=BLANC, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.8), Inches(7.7), Inches(14.4), Inches(0.5),
         "📸  [CAPTURE À INSÉRER — Console Airflow]",
         size=12, color=GRIS_MOYEN, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 15 — DOCKER & MLFLOW
# ============================================================
s = add_slide()
add_title(s, "Conteneurisation et suivi des expériences")

# Docker
add_rect(s, Inches(0.8), Inches(1.4), Inches(14.4), Inches(3.5),
         fill=GRIS_CLAIR, line=BLEU, radius=True)
add_rect(s, Inches(0.8), Inches(1.4), Inches(14.4), Inches(0.6),
         fill=BLEU, radius=True)
add_text(s, Inches(1.0), Inches(1.5), Inches(14.0), Inches(0.5),
         "🐳  DOCKER COMPOSE", size=18, bold=True, color=BLANC)

containers = [
    ("API", "Port 8000", BLEU_NUIT),
    ("AGENT", "Port 8001", BLEU_NUIT),
    ("MLFLOW", "Port 5000", BLEU_NUIT),
]
for i, (name, port, color) in enumerate(containers):
    x = Inches(1.2) + i * Inches(4.5)
    add_rect(s, x, Inches(2.3), Inches(4.0), Inches(1.2),
             fill=color, radius=True)
    add_text(s, x + Inches(0.2), Inches(2.45), Inches(3.6), Inches(0.5),
             name, size=20, bold=True, color=BLANC, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.95), Inches(3.6), Inches(0.4),
             port, size=13, color=TURQUOISE, align=PP_ALIGN.CENTER)

add_rect(s, Inches(1.2), Inches(3.7), Inches(13.6), Inches(1.0),
         fill=RGBColor(0xEC, 0xFD, 0xF5), line=VERT, radius=True)
add_text(s, Inches(1.4), Inches(3.85), Inches(13.2), Inches(0.4),
         "Volumes partagés", size=13, bold=True, color=VERT,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1.4), Inches(4.2), Inches(13.2), Inches(0.4),
         "/app/models (Read-Only)  •  /app/data", size=12,
         color=GRIS_FONCE, align=PP_ALIGN.CENTER)

# MLflow
add_rect(s, Inches(0.8), Inches(5.2), Inches(14.4), Inches(2.5),
         fill=GRIS_CLAIR, line=TURQUOISE, radius=True)
add_rect(s, Inches(0.8), Inches(5.2), Inches(14.4), Inches(0.6),
         fill=TURQUOISE, radius=True)
add_text(s, Inches(1.0), Inches(5.3), Inches(14.0), Inches(0.5),
         "📊  MLFLOW — Suivi des expériences", size=18, bold=True, color=BLANC)

mlflow_items = ["Enregistrement des runs", "Métriques (accuracy, F1, precision, recall)",
                "Paramètres des modèles", "Artefacts (modèles .pkl)"]
for i, item in enumerate(mlflow_items):
    add_text(s, Inches(1.2), Inches(6.1) + Inches(i * 0.4),
             Inches(13.8), Inches(0.4), f"•  {item}",
             size=13, color=GRIS_FONCE)

add_text(s, Inches(0.8), Inches(7.9), Inches(14.4), Inches(0.4),
         "⚠️  Streamlit n'est pas défini dans le Docker Compose observé",
         size=12, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 16 — CI/CD
# ============================================================
s = add_slide()
add_title(s, "Automatisation et CI/CD")

# Header
add_rect(s, Inches(3), Inches(1.3), Inches(10), Inches(0.8),
         fill=BLEU_NUIT, radius=True)
add_text(s, Inches(3.2), Inches(1.45), Inches(9.6), Inches(0.5),
         "GITHUB ACTIONS  —  .github/workflows/ci-cd.yml",
         size=15, bold=True, color=BLANC, align=PP_ALIGN.CENTER)

steps_cicd = [
    ("1. COMPILATION", "python -m compileall -q src tests", BLEU),
    ("2. TESTS", "pytest", TURQUOISE),
    ("3. VALIDATION DOCKER", "Validation de la configuration Docker Compose", BLEU),
    ("4. PUBLICATION", "Push des images sur GHCR", VERT),
]
for i, (title, desc, color) in enumerate(steps_cicd):
    y = Inches(2.4) + i * Inches(1.15)
    add_rect(s, Inches(2.5), y, Inches(11), Inches(0.95),
             fill=GRIS_CLAIR, line=color, radius=True)
    add_text(s, Inches(2.7), y + Inches(0.1), Inches(10.6), Inches(0.4),
             title, size=15, bold=True, color=color)
    add_text(s, Inches(2.7), y + Inches(0.5), Inches(10.6), Inches(0.4),
             desc, size=12, color=GRIS_FONCE)
    if i < len(steps_cicd) - 1:
        add_text(s, Inches(7.8), y + Inches(0.95), Inches(0.5), Inches(0.2),
                 "▼", size=14, bold=True, color=GRIS_MOYEN, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.8), Inches(7.5), Inches(14.4), Inches(0.5),
         "🔄  Déclenchement : push / pull request",
         size=15, bold=True, color=BLEU, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 17 — TESTS
# ============================================================
s = add_slide()
add_title(s, "Tests et validation technique")

# Tests présents
add_rect(s, Inches(0.8), Inches(1.4), Inches(9.5), Inches(3.0),
         fill=GRIS_CLAIR, line=BLEU, radius=True)
add_text(s, Inches(1.1), Inches(1.55), Inches(9.0), Inches(0.5),
         "✓  TESTS PRÉSENTS", size=16, bold=True, color=BLEU)
tests = [
    ("test_api_endpoints.py", "Endpoints REST"),
    ("test_agent.py", "Agent IA"),
    ("test_ticket_store.py", "Persistance SQLite"),
    ("test_environment_config.py", "Configuration"),
    ("test_predictor.py", "Prédicteur ML"),
    ("test_rag.py", "Recherche sémantique"),
]
for i, (name, desc) in enumerate(tests):
    y = Inches(2.1) + i * Inches(0.38)
    add_text(s, Inches(1.1), y, Inches(4.5), Inches(0.35),
             name, size=12, bold=True, color=GRIS_FONCE)
    add_text(s, Inches(5.8), y, Inches(4.2), Inches(0.35),
             f"→ {desc}", size=11, color=GRIS_MOYEN)

# Validation syntaxique
add_rect(s, Inches(10.6), Inches(1.4), Inches(4.6), Inches(3.0),
         fill=RGBColor(0xEC, 0xFD, 0xF5), line=VERT, radius=True)
add_text(s, Inches(10.8), Inches(1.55), Inches(4.2), Inches(0.5),
         "✓  VALIDATION SYNTAXIQUE", size=14, bold=True, color=VERT)
add_text(s, Inches(10.8), Inches(2.2), Inches(4.2), Inches(1.0),
         "python -m compileall\n-q src tests", size=11, color=GRIS_FONCE)
add_text(s, Inches(10.8), Inches(3.5), Inches(4.2), Inches(0.8),
         "✅  Compilation réussie\nsur l'archive extraite",
         size=12, bold=True, color=VERT)

# Distinction
add_rect(s, Inches(0.8), Inches(4.7), Inches(14.4), Inches(3.0),
         fill=RGBColor(0xFF, 0xFB, 0xEB), line=ORANGE, radius=True)
add_text(s, Inches(1.1), Inches(4.85), Inches(13.8), Inches(0.5),
         "⚠️  DISTINCTION IMPORTANTE", size=16, bold=True, color=ORANGE)
items = [
    "✓  Preuve de présence du code",
    "✓  Compilation réussie",
    "✓  Tests unitaires disponibles",
    "⚠️  Validation end-to-end non démontrée",
    "⚠️  Preuve de production non fournie",
]
for i, item in enumerate(items):
    y = Inches(5.4) + i * Inches(0.42)
    add_text(s, Inches(1.3), y, Inches(13.6), Inches(0.4),
             item, size=13, color=GRIS_FONCE)


# ============================================================
# SLIDE 18 — CONTRIBUTIONS
# ============================================================
s = add_slide()
add_title(s, "Contributions du projet")

# 3 blocs supérieurs
contribs_top = [
    ("DATA", "Data Engineering", "Snowflake • dbt\nDuckDB • SQLite", BLEU),
    ("ML", "Classification", "Type → Queue\n→ Priority", TURQUOISE),
    ("AIOps", "Prédiction\nexpérimentale", "MicroSS\nXGBoost", ORANGE),
]
for i, (title, subtitle, detail, color) in enumerate(contribs_top):
    x = Inches(0.8) + i * Inches(5.0)
    add_rect(s, x, Inches(1.4), Inches(4.5), Inches(3.0),
             fill=color, radius=True)
    add_text(s, x + Inches(0.2), Inches(1.6), Inches(4.1), Inches(0.7),
             title, size=28, bold=True, color=BLANC, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.5), Inches(4.1), Inches(0.5),
             subtitle, size=14, bold=True, color=BLANC, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(3.2), Inches(4.1), Inches(1.0),
             detail, size=12, color=BLANC, align=PP_ALIGN.CENTER)

# 2 blocs inférieurs
contribs_bottom = [
    ("RAG / IA", "Recherche sémantique + Agent IA\nFAISS + MiniLM • Human-in-the-loop", TURQUOISE),
    ("PLATEFORME", "FastAPI + Streamlit • Airflow + Docker + MLflow\nCI/CD GitHub Actions • 9 endpoints REST • 5 pages", BLEU),
]
for i, (title, detail, color) in enumerate(contribs_bottom):
    x = Inches(0.8) + i * Inches(7.4)
    add_rect(s, x, Inches(4.7), Inches(7.0), Inches(2.2),
             fill=color, radius=True)
    add_text(s, x + Inches(0.2), Inches(4.9), Inches(6.6), Inches(0.6),
             title, size=22, bold=True, color=BLANC, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(5.7), Inches(6.6), Inches(1.0),
             detail, size=12, color=BLANC, align=PP_ALIGN.CENTER)

# Synthèse
add_rect(s, Inches(0.8), Inches(7.1), Inches(14.4), Inches(1.4),
         fill=GRIS_CLAIR, line=TURQUOISE, radius=True)
add_text(s, Inches(1.1), Inches(7.3), Inches(13.8), Inches(1.0),
         "💡  Une base intégrée reliant Data Engineering, Machine Learning,\n"
         "     AIOps et assistance intelligente pour les opérations IT",
         size=14, bold=True, color=BLEU_NUIT, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 19 — LIMITES ET PERSPECTIVES
# ============================================================
s = add_slide()
add_title(s, "Limites actuelles et perspectives")

# Limites
add_rect(s, Inches(0.5), Inches(1.3), Inches(7.2), Inches(6.7),
         fill=RGBColor(0xFF, 0xFB, 0xEB), line=ORANGE, radius=True)
add_text(s, Inches(0.8), Inches(1.45), Inches(6.8), Inches(0.5),
         "⚠️  LIMITES", size=20, bold=True, color=ORANGE)

limites = [
    ("AIOps", "Performances faibles, nombreux faux positifs"),
    ("Données", "Historiques uniquement, pas de validation temps réel"),
    ("RAG", "Artefacts tickets.index et documents.pkl absents"),
    ("Validation", "End-to-end non démontrée"),
    ("Modèles", "Évaluation Queue/Priority à compléter"),
]
for i, (cat, detail) in enumerate(limites):
    y = Inches(2.2) + i * Inches(1.1)
    add_text(s, Inches(0.9), y, Inches(6.6), Inches(0.4),
             f"• {cat}", size=14, bold=True, color=ORANGE)
    add_text(s, Inches(1.2), y + Inches(0.4), Inches(6.3), Inches(0.5),
             detail, size=11, color=GRIS_FONCE)

# Perspectives
add_rect(s, Inches(8.3), Inches(1.3), Inches(7.2), Inches(6.7),
         fill=RGBColor(0xEF, 0xF6, 0xFF), line=BLEU, radius=True)
add_text(s, Inches(8.6), Inches(1.45), Inches(6.8), Inches(0.5),
         "🔮  PERSPECTIVES", size=20, bold=True, color=BLEU)

perspectives = [
    "Améliorer les modèles",
    "Réduire les faux positifs",
    "Améliorer le rappel",
    "Enrichir le corpus documentaire",
    "Rendre le RAG reproductible",
    "Monitoring temps réel",
    "Observabilité (Grafana)",
    "Suivi continu des modèles",
    "Streaming (Kafka) → PERSPECTIVE",
]
for i, item in enumerate(perspectives):
    y = Inches(2.2) + i * Inches(0.62)
    color = ORANGE if "Kafka" in item else GRIS_FONCE
    add_text(s, Inches(8.6), y, Inches(6.8), Inches(0.5),
             f"•  {item}", size=13, color=color)

add_text(s, Inches(0.5), Inches(8.1), Inches(15), Inches(0.5),
         "⚠️  Kafka = PERSPECTIVE UNIQUEMENT (non implémenté)",
         size=12, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 20 — CONCLUSION
# ============================================================
s = add_slide()
add_rect(s, 0, 0, prs.slide_width, prs.slide_height, fill=BLEU_NUIT)

add_text(s, Inches(0.5), Inches(0.5), Inches(15), Inches(0.7),
         "CONCLUSION", size=36, bold=True, color=BLANC, align=PP_ALIGN.CENTER)

# Chaîne DATA → INTELLIGENCE → ASSISTANCE
chain = [("DATA", TURQUOISE), ("INTELLIGENCE", TURQUOISE), ("ASSISTANCE", TURQUOISE)]
for i, (word, color) in enumerate(chain):
    x = Inches(2) + i * Inches(4.5)
    add_rect(s, x, Inches(1.6), Inches(3.5), Inches(1.0),
             fill=color, radius=True)
    add_text(s, x + Inches(0.2), Inches(1.75), Inches(3.1), Inches(0.7),
             word, size=20, bold=True, color=BLEU_NUIT, align=PP_ALIGN.CENTER)
    if i < 2:
        add_text(s, x + Inches(3.5), Inches(1.75), Inches(1.0), Inches(0.7),
                 "→", size=32, bold=True, color=BLANC, align=PP_ALIGN.CENTER)

# 3 contributions
contribs = [
    ("1. DATA", "Préparation et structuration des données"),
    ("2. INTELLIGENCE", "Classification des tickets et expérimentation AIOps"),
    ("3. ASSISTANCE", "Recherche sémantique, RAG et Agent IA via API / interface"),
]
for i, (title, desc) in enumerate(contribs):
    y = Inches(3.2) + i * Inches(1.1)
    add_rect(s, Inches(1.5), y, Inches(13), Inches(1.0),
             fill=RGBColor(0x1A, 0x2A, 0x42), line=TURQUOISE, radius=True)
    add_text(s, Inches(1.8), y + Inches(0.1), Inches(12.4), Inches(0.4),
             title, size=15, bold=True, color=TURQUOISE)
    add_text(s, Inches(1.8), y + Inches(0.5), Inches(12.4), Inches(0.5),
             desc, size=13, color=BLANC)

# Phrase finale
add_text(s, Inches(1), Inches(6.7), Inches(14), Inches(1.0),
         "💡  HARDTEC AIOps constitue une base d'expérimentation intégrée\n"
         "     rapprochant Data Engineering, Machine Learning et assistance\n"
         "     intelligente dans le domaine des opérations IT",
         size=14, color=BLANC, align=PP_ALIGN.CENTER)

add_text(s, Inches(1), Inches(8.0), Inches(14), Inches(0.5),
         "MERCI POUR VOTRE ATTENTION", size=22, bold=True,
         color=TURQUOISE, align=PP_ALIGN.CENTER)

add_text(s, Inches(1), Inches(8.5), Inches(14), Inches(0.4),
         "Questions ?", size=16, color=BLANC, align=PP_ALIGN.CENTER)


# ============================================================
# SAUVEGARDE
# ============================================================
output = "HARDTEC_AIOps_PFA_Presentation.pptx"
prs.save(output)
print(f"✅ Présentation générée : {output}")
print(f"📊 Nombre de slides : {len(prs.slides)}")